import asyncio
import io
import uuid
import zipfile
from datetime import UTC, date, datetime
from unittest.mock import Mock, patch

from openpyxl import Workbook, load_workbook

from app.api.v1.endpoints import macro_sap
from app.api.v1.endpoints.macro_sap import (
    _formatear_fecha_ultimo_gasto,
    _resaltar_uidd_necesario,
)


def test_formatear_fecha_ultimo_gasto_usa_la_fecha_mas_reciente() -> None:
    gastos = [
        {"spent_on": date(2026, 8, 10)},
        {"spent_on": date(2026, 8, 30)},
        {"spent_on": date(2026, 8, 18)},
    ]

    assert _formatear_fecha_ultimo_gasto(gastos) == "30/08/2026"


def test_resaltar_uidd_necesario_aplica_amarillo_en_columna_h() -> None:
    worksheet = Workbook().active
    worksheet.append(["S", "cuenta", 100, "V1", "T001", "", "CAJA CHICA", "DATO NECESARIO"])

    _resaltar_uidd_necesario(worksheet, worksheet.max_row, "DATO NECESARIO")

    assert worksheet["H1"].fill.fill_type == "solid"
    assert worksheet["H1"].fill.fgColor.rgb == "00FFFF00"
    assert worksheet["G1"].fill.fill_type is None


def test_resaltar_uidd_necesario_no_modifica_otro_uidd() -> None:
    worksheet = Workbook().active
    worksheet.append(["S", "cuenta", 100, "V1", "T001", "", "CAJA CHICA", "CFDI-123"])

    _resaltar_uidd_necesario(worksheet, worksheet.max_row, "CFDI-123")

    assert worksheet["H1"].fill.fill_type is None


def test_generar_polizas_conserva_lineas_separadas_con_misma_cuenta(tmp_path) -> None:
    solicitud_id = uuid.uuid4()
    solicitud = {
        "store_id": "store-1",
        "period_id": "period-1",
        "reimbursement_starts_on": date(2026, 8, 1),
        "reimbursement_ends_on": date(2026, 8, 31),
        "previous_reimbursement_starts_on": None,
        "previous_reimbursement_ends_on": None,
        "previous_reimbursement_amount": None,
        "created_at": datetime(2026, 9, 1, tzinfo=UTC),
        "folio": "TEST-1",
    }
    tienda = {"code": "V101"}
    gastos = [
        {
            "id": "expense-insumos",
            "category": "Insumos",
            "cfdi_uuid": "UUID-INSUMOS",
            "amount": "116",
            "cfdi_tax_rate": "16",
            "cfdi_tax_amount": "16",
            "cfdi_subtotal": "100",
            "spent_on": date(2026, 8, 10),
        },
        {
            "id": "expense-papeleria",
            "category": "Papelería",
            "cfdi_uuid": "UUID-PAPELERIA",
            "amount": "232",
            "cfdi_tax_rate": "16",
            "cfdi_tax_amount": "32",
            "cfdi_subtotal": "200",
            "spent_on": date(2026, 8, 11),
        },
    ]

    resultados = []
    for valor in (solicitud, tienda, {}, gastos):
        resultado = Mock()
        resultado.mappings.return_value.first.return_value = valor
        resultado.mappings.return_value.all.return_value = valor
        resultados.append(resultado)
    db = Mock()
    db.execute.side_effect = resultados

    plantilla = tmp_path / "plantilla.xlsx"
    Workbook().save(plantilla)

    with (
        patch.object(
            macro_sap,
            "diccionario_tiendas",
            {
                "V101": {
                    "PLAZA": "Tienda de prueba",
                    "NOMBRE_GERENTE": "Gerente",
                    "CUENTA": "Cuenta",
                    "CAJA_CHICA": "Fondo",
                    "RESPONSABLE": "Responsable",
                    "SUPERVISOR": "Supervisor",
                }
            },
        ),
        patch.object(
            macro_sap,
            "indice_categorias",
            {
                "insumos": {"codigo": "601007", "descripcion": "Insumos"},
                "papelería": {"codigo": "601007", "descripcion": "Papelería"},
            },
        ),
        patch.object(macro_sap, "tiendas_iva_w6", set()),
        patch.object(macro_sap, "ruta_plantilla", str(plantilla)),
        patch.object(macro_sap, "ruta_logo", str(tmp_path / "no-logo.png")),
        patch.object(
            macro_sap,
            "obtener_resumen_gasto_tienda",
            return_value={"current_accumulated": 0},
        ),
    ):
        response = macro_sap.generar_polizas(solicitud_id, db=db)

    async def leer_respuesta() -> bytes:
        return b"".join([parte async for parte in response.body_iterator])

    contenido_zip = asyncio.run(leer_respuesta())
    with zipfile.ZipFile(io.BytesIO(contenido_zip)) as archivo_zip:
        contenido_poliza = archivo_zip.read("Poliza Reembolso TEST-1.xlsx")

    poliza = load_workbook(io.BytesIO(contenido_poliza), data_only=True).active
    filas = {
        (poliza[f"C{fila}"].value, poliza[f"D{fila}"].value)
        for fila in range(41, poliza.max_row + 1)
        if poliza[f"C{fila}"].value is not None
    }

    assert ("601007", "Insumos") in filas
    assert ("601007", "Papelería") in filas
