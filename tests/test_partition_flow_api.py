"""Partition lifecycle checks against the real API and an isolated database."""

import asyncio
import io
import json
import os
import shutil
import socket
import subprocess
import threading
import time
import zipfile
from decimal import Decimal
from pathlib import Path
from uuid import UUID, uuid4

import httpx
import pytest
import uvicorn
from openpyxl import load_workbook
from sqlalchemy import select
from test_frontend_adapter_api import (
    _assign_user_to_store,
    _attach_valid_cfdi,
    _auth_headers,
    _create_user,
)

from app.api.v1.endpoints import macro_sap
from app.models.expense import Expense, ExpenseStatus
from app.models.reimbursement_request import ReimbursementRequest
from app.models.user import User


class PartitionFlow:
    def __init__(self, client, records, session_factory):
        self.client = client
        self.records = records
        self.sessions = session_factory
        self.request_id = records["request_id"]
        self.request_path = f"/api/v1/reimbursement-requests/{self.request_id}"
        self.frontend_path = f"/api/v1/frontend/solicitudes/{self.request_id}"
        self.actors = {}
        self.parent_id = None

    def actor(self, role):
        if role not in self.actors:
            email = f"partition-{role}@example.com"
            user_id = _create_user(self.client, role, email)
            _assign_user_to_store(self.client, self.records["store_id"], user_id, role)
            self.actors[role] = {
                "id": user_id,
                "email": email,
                "headers": _auth_headers(self.client, email),
            }
        return self.actors[role]

    def move(self, role, target):
        response = self.client.post(
            f"{self.request_path}/transition/me",
            headers=self.actor(role)["headers"],
            json={"target_status": target, "note": f"Partition flow: {target}"},
        )
        assert response.status_code == 200, response.text
        return response.json()

    def capture(self, *, authorization=False):
        response = self.client.post(
            f"{self.frontend_path}/gastos/me",
            headers=self.actor("store")["headers"],
            json={
                "fecha": "07/08/2026",
                "categoria": "Pasajes y Taxis" if authorization else "Agua",
                "monto": "1500.00",
                "observaciones": "Nota original de tienda",
                **({"areaAutoriza": "Supervisores"} if authorization else {}),
            },
        )
        assert response.status_code == 201, response.text
        self.parent_id = response.json()["gastos"][0]["backendId"]
        _attach_valid_cfdi(self.client, self.parent_id, "1500.00", uuid=str(uuid4()))
        return self.detail()

    def accounting(self, *, authorization=False):
        self.capture(authorization=authorization)
        self.move("store", "submitted")
        if authorization:
            authorizer = self.actor("authorizer")
            self.move("authorizer", "authorization_review")
            # A pending authorization must not enter accounting, even with a positive amount.
            blocked = self.client.post(
                f"{self.request_path}/transition/me",
                headers=self.actor("accountant")["headers"],
                json={"target_status": "under_accounting_review"},
            )
            assert blocked.status_code == 409, blocked.text
            authorized = self.client.post(
                f"/api/v1/expenses/{self.parent_id}/authorize/me",
                headers=authorizer["headers"],
                json={"note": "Autorizado por supervisor de tienda"},
            )
            assert authorized.status_code == 200, authorized.text
            self.move("authorizer", "authorized")
        self.move("accountant", "under_accounting_review")

    def detail(self, role="accountant"):
        response = self.client.get(f"{self.frontend_path}/me", headers=self.actor(role)["headers"])
        assert response.status_code == 200, response.text
        return response.json()

    def split_path(self, expense_id=None):
        return f"{self.frontend_path}/gastos/{expense_id or self.parent_id}/particiones/me"

    def split(self, parts=None, *, role="accountant", expense_id=None):
        return self.client.post(
            self.split_path(expense_id),
            headers=self.actor(role)["headers"],
            json={"particiones": parts or self.parts()},
        )

    @staticmethod
    def parts():
        return [
            {"categoria": "Papelería", "monto": "1160.00", "impuesto": "16"},
            {"categoria": "Agua", "monto": "340.00", "impuesto": "0"},
        ]

    def children(self, payload=None):
        return sorted(
            [g for g in (payload or self.detail())["gastos"] if g["idOriginal"] == self.parent_id],
            key=lambda g: g["particionIndex"],
        )

    def events(self, *, observations=False):
        response = self.client.get(
            f"{self.request_path}/audit-events",
            headers=self.actor("accountant")["headers"],
            params={"observations_only": observations, "limit": 200},
        )
        assert response.status_code == 200, response.text
        return response.json()

    def note(self, expense_id, text):
        response = self.client.post(
            f"/api/v1/expenses/{expense_id}/observation/me",
            headers=self.actor("accountant")["headers"],
            json={"note": text},
        )
        assert response.status_code == 200, response.text


@pytest.fixture
def partition_flow(client, base_records, session_factory):
    return PartitionFlow(client, base_records, session_factory)


def test_partition_full_workflow_through_payment(partition_flow, test_settings):
    flow = partition_flow
    assert not test_settings.textract_enabled and not test_settings.ses_enabled
    flow.accounting(authorization=True)
    original = flow.detail()["gastos"][0]
    split = flow.split()
    assert split.status_code == 200, split.text
    children = flow.children()
    assert len(children) == 2
    assert flow.detail()["expenseCount"] == 2
    assert flow.detail()["montoTotal"] == 1500
    for child in children:
        assert child["folioFiscal"] == original["folioFiscal"]
        assert child["urlFactura"] == original["urlFactura"]
        assert child["autorizacion"] == "autorizado"
        downloaded = flow.client.get(
            child["urlFactura"], headers=flow.actor("accountant")["headers"]
        )
        assert downloaded.status_code == 200
        assert original["folioFiscal"].encode() in downloaded.content.upper()
    assert children[0]["cfdiSubtotal"] == 1000
    assert children[0]["cfdiTaxAmount"] == 160
    assert children[0]["sapTaxIndexOverride"] == "W1"
    assert children[1]["cfdiTaxRate"] == 0
    assert children[1]["cfdiSubtotal"] == 340
    summary = flow.client.get(f"{flow.request_path}/validation-summary").json()
    assert summary["missing_authorization_expense_ids"] == []
    assert summary["missing_cfdi_expense_ids"] == []
    assert summary["duplicate_cfdi_uuids"] == []
    flow.note(children[0]["backendId"], "Nota de la primera partición")
    flow.move("accountant", "accounting_reviewed")
    prepared = flow.client.post(
        f"{flow.request_path}/sap-policy/prepare/me",
        headers=flow.actor("accountant")["headers"],
        json={"reference": "SAP-PARTITION-FLOW", "note": "Exportación de prueba"},
    )
    assert prepared.status_code == 200, prepared.text
    flow.move("accountant", "accounting_manager_review")
    flow.move("accounting_manager", "accounting_manager_approved")
    flow.move("treasury", "treasury_review")
    flow.move("treasury", "direction_approved")
    for role in ["accounting_manager", "treasury", "director"]:
        detail = flow.detail(role)
        assert detail["expenseCount"] == 2 and detail["montoTotal"] == 1500
        assert {g["backendId"] for g in flow.children(detail)} == {g["backendId"] for g in children}
        blocked = flow.split(role=role)
        assert blocked.status_code == 403, blocked.text
    flow.move("accounting_manager", "approved_for_payment")
    payment = flow.client.post(
        f"{flow.request_path}/payments/me",
        headers=flow.actor("accounting_manager")["headers"],
        json={"reference": "PARTITION-TEST-PAYMENT", "note": "Prueba aislada; sin pago real"},
    )
    assert payment.status_code == 201, payment.text
    assert flow.detail()["backendStatus"] == "paid"
    assert flow.split().status_code == 409
    assert (
        flow.client.delete(
            flow.split_path(), headers=flow.actor("accountant")["headers"]
        ).status_code
        == 409
    )
    locked_note = flow.client.post(
        f"/api/v1/expenses/{children[0]['backendId']}/observation/me",
        headers=flow.actor("accountant")["headers"],
        json={"note": "Bloqueado después de pago"},
    )
    assert locked_note.status_code == 409
    assert any(
        e["message"] == "Nota de la primera partición" for e in flow.events(observations=True)
    )


def test_partition_edit_cancel_repartition_keeps_only_current_children(partition_flow):
    flow = partition_flow
    flow.accounting()
    assert flow.split().status_code == 200
    first_ids = {c["backendId"] for c in flow.children()}
    replacement = [
        {"categoria": "Servicio de Agua", "monto": "1080", "impuesto": "8"},
        {"categoria": "Bolsas", "monto": "420", "impuesto": "0"},
    ]
    edited = flow.split(replacement)
    assert edited.status_code == 200, edited.text
    second = flow.children()
    assert {c["backendId"] for c in second} == first_ids
    assert second[0]["cfdiSubtotal"] == 1000 and second[0]["cfdiTaxAmount"] == 80
    assert second[0]["sapTaxIndexOverride"] == "W6"
    cancel = flow.client.delete(
        flow.split_path(second[0]["backendId"]), headers=flow.actor("accountant")["headers"]
    )
    assert cancel.status_code == 200, cancel.text
    restored = flow.detail()
    assert restored["expenseCount"] == 1 and restored["montoTotal"] == 1500
    assert flow.children(restored) == []
    assert restored["gastos"][0]["inactivo"] is False
    repeated = flow.client.delete(flow.split_path(), headers=flow.actor("accountant")["headers"])
    assert repeated.status_code == 409
    assert flow.split().status_code == 200
    assert {c["backendId"] for c in flow.children()}.isdisjoint(first_ids)
    assert flow.detail()["expenseCount"] == 2
    events = flow.events(observations=True)
    assert sum(e["action"] == "expense_partitioned" for e in events) == 2
    changed = next(e for e in events if e["action"] == "expense_partition_updated")
    assert "categoría de Papelería a Servicio de Agua" in changed["message"]
    assert "monto de $1,160.00 a $1,080.00" in changed["message"]
    assert "impuesto de 16% a 8%" in changed["message"]
    assert any(e["action"] == "expense_partition_cancelled" for e in events)
    with flow.sessions() as db:
        rows = list(
            db.scalars(
                select(Expense).where(Expense.reimbursement_request_id == UUID(flow.request_id))
            )
        )
        assert len(rows) == 5
        assert sum(row.status == ExpenseStatus.removed for row in rows) == 2


@pytest.mark.parametrize(
    "invalid_parts",
    [
        [{"categoria": "Agua", "monto": "1500"}],
        [{"categoria": "Agua", "monto": "1000"}, {"categoria": "Bolsas", "monto": "499"}],
        [{"categoria": "Agua", "monto": "1500"}, {"categoria": "Bolsas", "monto": "0"}],
        [{"categoria": "Agua", "monto": "1501"}, {"categoria": "Bolsas", "monto": "-1"}],
        [{"categoria": "", "monto": "1000"}, {"categoria": "Bolsas", "monto": "500"}],
        [
            {"categoria": "Agua", "monto": "1000", "impuesto": "-1"},
            {"categoria": "Bolsas", "monto": "500"},
        ],
    ],
)
def test_invalid_partition_is_atomic(partition_flow, invalid_parts):
    flow = partition_flow
    flow.accounting()
    before = flow.detail()
    events_before = len(flow.events())
    result = flow.split(invalid_parts)
    assert result.status_code == 422, result.text
    assert flow.detail() == before
    assert len(flow.events()) == events_before


def test_partition_role_stage_identity_and_nested_guards(partition_flow):
    flow = partition_flow
    flow.capture()
    assert flow.split().status_code == 409
    flow.move("store", "submitted")
    assert flow.split().status_code == 409
    flow.move("accountant", "under_accounting_review")
    assert (
        flow.client.post(flow.split_path(), json={"particiones": flow.parts()}).status_code == 401
    )
    for role in ["store", "authorizer", "accounting_manager", "treasury", "director"]:
        assert flow.split(role=role).status_code == 403
    assert flow.split(expense_id=str(uuid4())).status_code == 404
    assert flow.split().status_code == 200
    child_id = flow.children()[0]["backendId"]
    nested = flow.split(expense_id=child_id)
    assert nested.status_code == 409
    assert nested.json()["detail"]["code"] == "PARTITION_CHILD_CANNOT_BE_PARTITIONED"
    other_store = flow.client.post(
        "/api/v1/stores/", json={"code": "OTHER", "name": "Otra tienda"}
    ).json()
    other_request = flow.client.post(
        "/api/v1/reimbursement-requests/",
        json={
            "store_id": other_store["id"],
            "period_id": flow.records["period_id"],
            "reported_total": "1",
        },
    ).json()
    path = (
        f"/api/v1/frontend/solicitudes/{other_request['id']}/gastos/{flow.parent_id}/particiones/me"
    )
    result = flow.client.post(
        path, headers=flow.actor("accountant")["headers"], json={"particiones": flow.parts()}
    )
    assert result.status_code in {404, 409}, result.text
    assert flow.detail()["montoTotal"] == 1500


def test_partition_permissions_match_all_roles_and_workflow_stages(partition_flow):
    from app.models.reimbursement_request import ReimbursementRequestStatus

    flow = partition_flow
    flow.accounting()
    assert flow.split().status_code == 200
    roles = [
        "store",
        "authorizer",
        "accountant",
        "accounting_manager",
        "treasury",
        "director",
        "admin",
    ]
    for role in roles:
        flow.actor(role)
    for stage in ReimbursementRequestStatus:
        for role in roles:
            with flow.sessions() as db:
                request = db.get(ReimbursementRequest, UUID(flow.request_id))
                request.status = stage
                db.commit()
            allowed_role = role in {"accountant", "admin"}
            allowed_stage = stage.value in {"under_accounting_review", "accounting_reviewed"}
            expected = 200 if allowed_role and allowed_stage else 409 if allowed_role else 403
            posted = flow.split(role=role)
            assert posted.status_code == expected, (role, stage.value, posted.text)
            deleted = flow.client.delete(flow.split_path(), headers=flow.actor(role)["headers"])
            assert deleted.status_code == expected, (role, stage.value, deleted.text)
            if expected == 200:
                assert flow.split(role=role).status_code == 200


def test_removing_every_partition_does_not_restore_the_original_amount(partition_flow):
    flow = partition_flow
    flow.accounting()
    assert flow.split().status_code == 200
    for child in flow.children():
        response = flow.client.post(
            f"/api/v1/expenses/{child['backendId']}/remove/me",
            headers=flow.actor("accountant")["headers"],
            json={"reason": "No procede esta parte", "adjust_reported_total": True},
        )
        assert response.status_code == 200, response.text
    detail = flow.detail()
    assert detail["montoTotal"] == 0, (
        "Removing all children must not silently restore the original invoice"
    )
    assert detail["expenseCount"] == 0


def test_partition_edit_keeps_previous_child_note_visible_on_current_child(partition_flow):
    flow = partition_flow
    flow.accounting()
    assert flow.split().status_code == 200
    previous_child = flow.children()[0]
    flow.note(previous_child["backendId"], "Nota que debe seguir visible al editar")
    updated = flow.parts()
    updated[0]["impuesto"] = "8"
    assert flow.split(updated).status_code == 200
    current_child = flow.children()[0]
    history = flow.events(observations=True)
    note = next(e for e in history if e["message"] == "Nota que debe seguir visible al editar")
    # These are exactly the IDs selected by historialParaGasto in the current UI.
    visible_ids = {current_child["backendId"], current_child["idOriginal"]}
    assert note["expense_id"] in visible_ids, "The saved note is stranded on the replaced child ID"


def test_comments_survive_partition_count_changes_cancellation_and_recreation(partition_flow):
    flow = partition_flow
    flow.accounting()
    parts = [{"categoria": "Agua", "monto": "500", "impuesto": "0"} for _ in range(3)]
    assert flow.split(parts).status_code == 200
    first = flow.children()
    for index, child in enumerate(first):
        flow.note(child["backendId"], f"Comentario de parte {index + 1}")
    assert flow.split().status_code == 200
    current = flow.children()
    assert [c["backendId"] for c in current] == [c["backendId"] for c in first[:2]]
    notes = [e for e in flow.events(observations=True) if e["action"] == "expense_observation_added"]
    removed_note = next(e for e in notes if e["message"] == "Comentario de parte 3")
    assert removed_note["expense_archived"] is True
    assert removed_note["expense_parent_id"] == flow.parent_id
    assert removed_note["expense_partition_index"] == 3
    assert flow.client.delete(flow.split_path(), headers=flow.actor("accountant")["headers"]).status_code == 200
    assert flow.split().status_code == 200
    flow.note(flow.children()[0]["backendId"], "Comentario de la partición nueva")
    notes = [e for e in flow.events(observations=True) if e["action"] == "expense_observation_added"]
    assert len(notes) == 5  # Three previous comments, the new one, and the store's initial note.
    assert len({e["id"] for e in notes}) == 5
    assert sum(e["expense_archived"] for e in notes) == 3
    assert flow.detail()["montoTotal"] == 1500


def test_partition_edit_invalidates_an_already_prepared_sap_policy(partition_flow):
    flow = partition_flow
    flow.accounting()
    assert flow.split().status_code == 200
    flow.move("accountant", "accounting_reviewed")
    prepared = flow.client.post(
        f"{flow.request_path}/sap-policy/prepare/me",
        headers=flow.actor("accountant")["headers"],
        json={"reference": "BEFORE-PARTITION-EDIT"},
    )
    assert prepared.status_code == 200, prepared.text
    parts = flow.parts()
    parts[0]["categoria"] = "Servicio de Agua"
    assert flow.split(parts).status_code == 200
    advanced = flow.client.post(
        f"{flow.request_path}/transition/me",
        headers=flow.actor("accountant")["headers"],
        json={"target_status": "accounting_manager_review"},
    )
    assert advanced.status_code == 409, (
        "A stale policy must not allow advancement after a partition edit"
    )
    with flow.sessions() as db:
        request = db.get(ReimbursementRequest, UUID(flow.request_id))
        assert request.sap_policy_generated_at is None, (
            "The previously exported policy no longer matches the partition"
        )


class _TypedSqliteMacroSession:
    """Keep real SQL/results while supplying the typed values PostgreSQL returns."""

    def __init__(self, session):
        self.session = session

    def __getattr__(self, name):
        return getattr(self.session, name)

    def execute(self, statement, params):
        from app.models.period import Period
        from app.models.store import Store

        sql = str(statement)
        if self.session.bind.dialect.name == "sqlite":
            params = {
                key: value.hex if isinstance(value, UUID) else value
                for key, value in params.items()
            }
            for table in [
                Expense.__table__,
                ReimbursementRequest.__table__,
                Store.__table__,
                Period.__table__,
            ]:
                if f"FROM {table.name}" in sql:
                    statement = statement.columns(**{col.name: col.type for col in table.columns})
                    break
        return self.session.execute(statement, params)


def test_macro_exports_only_active_partitions_with_shared_folio_and_manual_taxes(
    partition_flow, test_settings
):
    flow = partition_flow
    flow.accounting()
    assert flow.split().status_code == 200
    assert (
        flow.client.delete(
            flow.split_path(), headers=flow.actor("accountant")["headers"]
        ).status_code
        == 200
    )
    assert (
        flow.split(
            [
                {"categoria": "Papelería", "monto": "1080", "impuesto": "8"},
                {"categoria": "Agua", "monto": "420", "impuesto": "0"},
            ]
        ).status_code
        == 200
    )
    children = flow.children()
    with flow.sessions() as db:
        accountant = db.get(User, UUID(flow.actor("accountant")["id"]))
        response = macro_sap.generar_polizas(
            UUID(flow.request_id), db=_TypedSqliteMacroSession(db), current_user=accountant
        )

        async def read_body():
            return b"".join([part async for part in response.body_iterator])

        content = asyncio.run(read_body())
    with zipfile.ZipFile(io.BytesIO(content)) as archive:
        sap_name = next(name for name in archive.namelist() if name.startswith("CAJA CHICA"))
        sap = load_workbook(io.BytesIO(archive.read(sap_name)), data_only=True).active
    rows = list(sap.iter_rows(values_only=True))
    debit = [row for row in rows if row[0] == "S"]
    credit = [row for row in rows if row[0] == "K"]
    assert len(debit) == 2
    assert {Decimal(str(row[2])) for row in debit} == {Decimal(1080), Decimal(420)}
    assert {row[3] for row in debit} == {"W6", "W0"}
    assert {row[7] for row in debit} == {children[0]["folioFiscal"]}
    assert len(credit) == 1 and Decimal(str(credit[0][2])) == Decimal(-1500)
    # Upload the actual generated workbook through the same API used by accounting.
    with zipfile.ZipFile(io.BytesIO(content)) as archive:
        workbook_name = next(
            name for name in archive.namelist() if name.startswith("Poliza Reembolso")
        )
        workbook_bytes = archive.read(workbook_name)
    test_settings.max_upload_bytes = 1024 * 1024
    uploaded = flow.client.post(
        f"{flow.request_path}/reimbursement-excel/me",
        headers=flow.actor("accountant")["headers"],
        files={
            "file": (
                workbook_name,
                workbook_bytes,
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
    )
    assert uploaded.status_code == 201, uploaded.text
    detail = flow.detail()
    downloaded = flow.client.get(
        detail["reembolsoDownloadUrl"], headers=flow.actor("accountant")["headers"]
    )
    assert downloaded.status_code == 200
    assert downloaded.content == workbook_bytes


@pytest.mark.skipif(
    os.getenv("SMOLBOX_PARTITION_BROWSER") != "1", reason="Opt-in real browser/server integration"
)
def test_real_browser_partition_navigation(partition_flow, test_app, tmp_path):
    flow = partition_flow
    flow.accounting()
    root = Path(__file__).resolve().parents[1]
    node = os.getenv("NODE_EXECUTABLE") or shutil.which("node")
    assert node, "Set NODE_EXECUTABLE to the installed Node runtime"
    # Bind only loopback and let the OS choose unused ports; never reuse a live app server.
    api_socket = socket.socket()
    api_socket.bind(("127.0.0.1", 0))
    api_socket.listen(128)
    api_port = api_socket.getsockname()[1]
    with socket.socket() as frontend_socket:
        frontend_socket.bind(("127.0.0.1", 0))
        frontend_port = frontend_socket.getsockname()[1]
    server = uvicorn.Server(uvicorn.Config(test_app, log_level="warning"))
    thread = threading.Thread(target=server.run, kwargs={"sockets": [api_socket]}, daemon=True)
    thread.start()
    env = {
        **os.environ,
        "VITE_COGNITO_ENABLED": "false",
        "PARTITION_TEST_API_URL": f"http://127.0.0.1:{api_port}",
        "FRONTEND_TEST_URL": f"http://127.0.0.1:{frontend_port}",
        "PARTITION_TEST_DATA": json.dumps(
            {
                "requestId": flow.request_id,
                "parentId": flow.parent_id,
                "folio": flow.detail()["folio"],
                "email": flow.actor("accountant")["email"],
                "outputDir": str(tmp_path),
            }
        ),
    }
    try:
        with (tmp_path / "vite.log").open("w") as log:
            vite = subprocess.Popen(
                [
                    node,
                    "node_modules/vite/bin/vite.js",
                    "--config",
                    "tests/partitionFlow.vite.mjs",
                    "--host",
                    "127.0.0.1",
                    "--port",
                    str(frontend_port),
                    "--strictPort",
                ],
                cwd=root / "frontend",
                env=env,
                stdout=log,
                stderr=subprocess.STDOUT,
            )
            try:
                for _ in range(150):
                    try:
                        if httpx.get(env["FRONTEND_TEST_URL"], timeout=1).status_code == 200:
                            break
                    except httpx.HTTPError:
                        pass
                    time.sleep(0.1)
                else:
                    pytest.fail("The isolated frontend server did not start")
                scripts = ["partitionFlow.browser.mjs"]
                if os.getenv("SMOLBOX_BROWSER_REGRESSIONS") == "1":
                    scripts.extend(["observationHistory.browser.mjs", "authorizerNavigation.browser.mjs"])
                for script in scripts:
                    result = subprocess.run(
                        [node, f"tests/{script}"], cwd=root / "frontend", env=env,
                        text=True, capture_output=True, timeout=180, check=False,
                    )
                    print(result.stdout)
                    assert result.returncode == 0, script + "\n" + result.stdout + result.stderr
            finally:
                vite.terminate()
                vite.wait(timeout=15)
    finally:
        server.should_exit = True
        thread.join(timeout=15)
        api_socket.close()
    assert not thread.is_alive()
